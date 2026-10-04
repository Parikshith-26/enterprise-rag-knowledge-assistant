from pathlib import Path
import re

from docx import Document as DocxDocument

from .base import Document


def load_docx(file_path: str, company: str) -> Document:
    path = Path(file_path)

    docx = DocxDocument(path)

    paragraphs = []

    for paragraph in docx.paragraphs:
        text = paragraph.text.strip()

        if not text:
            continue

        # Preserve Word heading structure
        style_name = ""

        if paragraph.style:
            style_name = paragraph.style.name.lower()

        if style_name.startswith("heading"):
            match = re.search(r"(\d+)", style_name)

            if match:
                level = min(int(match.group(1)), 6)
            else:
                level = 2

            paragraphs.append(
                f"{'#' * level} {text}"
            )

        else:
            paragraphs.append(text)

    full_text = "\n\n".join(paragraphs)

    return Document(
        document_id=path.stem,
        source=str(path),
        title=path.stem,
        document_type="docx",
        text=full_text,
        metadata={
            "company": company,
            "file_name": path.name,
        },
    )