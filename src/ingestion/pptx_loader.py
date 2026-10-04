from pathlib import Path
import re

from pptx import Presentation

from .base import Document


def is_likely_heading(text: str) -> bool:
    """
    Detect common heading patterns in PowerPoint text.
    """

    text = text.strip()

    if not text:
        return False

    # Numbered headings
    if re.match(r"^\d+[\.\)]\s+\S+", text):
        return True

    # Numbered sub-headings
    if re.match(r"^\d+\.\d+\s+\S+", text):
        return True

    heading_keywords = [
        "Key Features",
        "Who Can Apply?",
        "Documents Required",
        "Eligibility",
        "How to Apply",
        "Application Process",
        "Benefits",
        "Features",
        "Important Information",
        "Supporting Rural Dreams",
    ]

    for keyword in heading_keywords:
        if text.lower() == keyword.lower():
            return True

    return False


def load_pptx(file_path: str, company: str) -> Document:
    path = Path(file_path)

    presentation = Presentation(path)

    slides = []

    for slide_number, slide in enumerate(
        presentation.slides,
        start=1
    ):

        slide_text = []

        # First try to identify the slide title.
        title_shape = slide.shapes.title

        title_text = ""

        if title_shape is not None:
            title_text = title_shape.text.strip()

        if title_text:
            slide_text.append(title_text)

        # Extract remaining shapes.
        for shape in slide.shapes:

            if shape == title_shape:
                continue

            if not hasattr(shape, "text"):
                continue

            text = shape.text.strip()

            if not text:
                continue

            for line in text.splitlines():

                line = line.strip()

                if line:
                    slide_text.append(line)

        if slide_text:

            slides.append(
                f"[Slide {slide_number}]\n"
                + "\n".join(slide_text)
            )

    full_text = "\n\n".join(slides)

    return Document(
        document_id=path.stem,
        source=str(path),
        title=path.stem,
        document_type="pptx",
        text=full_text,
        metadata={
            "company": company,
            "file_name": path.name,
            "slide_count": len(presentation.slides),
        },
    )